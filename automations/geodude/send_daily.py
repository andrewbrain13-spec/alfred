#!/usr/bin/env python3
"""Geodude daily email: when a new day's run lands in the folder, email it.

Triggered by files appearing (not a fixed time): Alfred runs this every few
minutes. When it sees a new `Geodude_Summary_<YYYYMMDD>.docx`, it waits a
settle period (so all the day's files finish landing), then emails the Summary
doc as the body with every file in that batch attached, and records the batch
so it never sends twice.

Env:
  GEODUDE_FOLDER          folder to watch (default: the OneDrive daily-runs path)
  GEODUDE_TO              comma list (default: abrain@braingroup.com)
  GEODUDE_CC              comma list (default: chad@chadsneed.com, jorscheln@braingroup.com)
  GEODUDE_SETTLE_SECONDS  wait after the summary is first seen (default 600 = 10 min)
  GEODUDE_STATE_FILE      state file (default geodude_state.json)
  GEODUDE_DRY_RUN         "1" = log only, don't send (default "1")
  ALFRED_TENANT_ID / ALFRED_CLIENT_ID   Graph app registration (Mail.Send required)
"""

from __future__ import annotations

import json
import os
import re
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from alfred.graph import GraphClient

DEFAULT_FOLDER = (
    r"C:\Users\abrai\OneDrive - Brain Group\Documents - Brain Group Shared Drive"
    r"\5. Investments\Oil Royalties\Geodude daily runs"
)
_DATE_RE = re.compile(r"(\d{8})\.[^.]+$")
_MONTHS = ["", "January", "February", "March", "April", "May", "June", "July",
           "August", "September", "October", "November", "December"]


def extract_date_token(filename: str) -> str:
    """The YYYYMMDD token before the extension, or '' if absent."""
    m = _DATE_RE.search(filename)
    return m.group(1) if m else ""


def group_by_date(filenames: list[str]) -> dict[str, list[str]]:
    """Group Geodude_* files by their date token."""
    groups: dict[str, list[str]] = {}
    for f in filenames:
        if not f.lower().startswith("geodude"):
            continue
        tok = extract_date_token(f)
        if tok:
            groups.setdefault(tok, []).append(f)
    return groups


def format_subject_date(token: str) -> str:
    """20260710 -> 'July 10, 2026' (fallback when the doc has no title)."""
    try:
        y, m, d = int(token[:4]), int(token[4:6]), int(token[6:8])
        return f"{_MONTHS[m]} {d}, {y}"
    except (ValueError, IndexError):
        return token


def _summary_name(names: list[str]) -> str | None:
    for n in names:
        if "summary" in n.lower() and n.lower().endswith(".docx"):
            return n
    return None


def _subject(summary_path: str, token: str) -> str:
    """Use the doc's Title paragraph if present, else build from the date."""
    try:
        import docx

        d = docx.Document(summary_path)
        for p in d.paragraphs:
            if p.style and p.style.name == "Title" and p.text.strip():
                return p.text.strip()
    except Exception:
        pass
    return f"EnergyNet Daily Royalty Screen — {format_subject_date(token)}"


def _body_html(summary_path: str) -> str:
    """Convert the Summary .docx to HTML for the email body."""
    try:
        import mammoth

        with open(summary_path, "rb") as fh:
            html = mammoth.convert_to_html(fh).value
        if html.strip():
            return html
    except Exception:
        pass
    # Fallback: plain-text paragraphs wrapped as simple HTML.
    try:
        import docx

        d = docx.Document(summary_path)
        paras = [p.text for p in d.paragraphs if p.text.strip()]
        return "<br>".join(_escape(p) for p in paras)
    except Exception:
        return "See attached files."


def _escape(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _load_state(path: str) -> dict:
    if os.path.exists(path):
        try:
            data = json.load(open(path, encoding="utf-8"))
            data.setdefault("processed", [])
            data.setdefault("first_seen", {})
            return data
        except Exception:
            pass
    return {"processed": [], "first_seen": {}}


def _save_state(path: str, state: dict) -> None:
    tmp = f"{path}.tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(state, fh)
    os.replace(tmp, path)


def main() -> int:
    folder = os.environ.get("GEODUDE_FOLDER", DEFAULT_FOLDER)
    to = [a.strip() for a in os.environ.get("GEODUDE_TO", "abrain@braingroup.com").split(",") if a.strip()]
    cc = [a.strip() for a in os.environ.get(
        "GEODUDE_CC", "chad@chadsneed.com,jorscheln@braingroup.com").split(",") if a.strip()]
    settle = int(os.environ.get("GEODUDE_SETTLE_SECONDS", "600"))
    state_file = os.environ.get("GEODUDE_STATE_FILE", "geodude_state.json")
    dry_run = os.environ.get("GEODUDE_DRY_RUN", "1") == "1"

    if not os.path.isdir(folder):
        print(f"Geodude: folder not found: {folder}")
        return 0

    groups = group_by_date([f for f in os.listdir(folder) if f.lower().startswith("geodude")])
    state = _load_state(state_file)
    now = time.time()

    for token in sorted(groups):
        if token in state["processed"]:
            continue
        names = groups[token]
        summary = _summary_name(names)
        if not summary:
            continue  # wait until the summary doc appears

        first = state["first_seen"].setdefault(token, now)
        _save_state(state_file, state)
        waited = now - first
        if waited < settle:
            print(f"Geodude batch {token}: summary present, settling "
                  f"({int(settle - waited)}s left before sending).")
            continue

        summary_path = os.path.join(folder, summary)
        subject = _subject(summary_path, token)
        body = _body_html(summary_path)
        attachments = []
        for n in sorted(names):
            with open(os.path.join(folder, n), "rb") as fh:
                attachments.append((n, fh.read()))

        if dry_run:
            print(f"DRY-RUN: would send {subject!r} to {to} cc {cc} with "
                  f"{len(attachments)} attachment(s) (batch {token}).")
        else:
            GraphClient().send_mail(subject, body, to, cc, attachments)
            print(f"SENT {subject!r} to {to} cc {cc}, {len(attachments)} attachment(s) (batch {token}).")

        state["processed"].append(token)
        _save_state(state_file, state)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
